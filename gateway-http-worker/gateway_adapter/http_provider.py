from __future__ import annotations

import asyncio
import base64
import csv
import hashlib
import io
import json
import math
import random
import re
import time
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from html.parser import HTMLParser
from typing import Any
from urllib.parse import urlencode, urljoin, urlparse

from curl_cffi.requests import AsyncSession
from curl_cffi.requests.exceptions import ConnectionError as CurlConnectionError
from curl_cffi.requests.exceptions import Timeout as CurlTimeout
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


REQUIRED_COLUMNS = (
    "File Date",
    "Producer",
    "Producer NPN",
    "Email",
    "Upline",
    "Upline NPN",
    "Agency",
    "All Lines Personal",
    "All Lines Team",
    "1st Leg Producer",
    "1st Leg NPN",
    "1st Leg %",
    "Outside 1st Leg",
)

PERIOD_CODES = {"wtd": "1", "mtd": "3", "ytd": "5"}
POLL_SCHEDULE_SECONDS = (2.0, 3.0, 5.0, 8.0, 10.0)


class GatewayContractError(RuntimeError):
    """The private Gateway/StatS web contract no longer has the expected shape."""


class GatewayAuthenticationError(RuntimeError):
    """Credentials/session cannot establish the configured upstream account."""


class GatewayTransientError(RuntimeError):
    """The upstream operation failed in a way that is safe to retry."""


@dataclass(frozen=True, slots=True)
class HttpGatewayConfig:
    gateway_origin: str = "https://ww3.familyfirstlife.com"
    stats_origin: str = "https://stats.uppatop.com"
    request_timeout_seconds: float = 30.0
    export_timeout_seconds: float = 120.0
    ytd_timeout_seconds: float = 180.0
    impersonate: str = "chrome"

    def __post_init__(self) -> None:
        for label, value in (("gateway_origin", self.gateway_origin), ("stats_origin", self.stats_origin)):
            parsed = urlparse(value)
            if (
                parsed.scheme != "https"
                or not parsed.netloc
                or parsed.username is not None
                or parsed.password is not None
                or parsed.path not in ("", "/")
                or parsed.params
                or parsed.query
                or parsed.fragment
            ):
                raise GatewayContractError(f"{label} must be a bare trusted HTTPS origin.")


@dataclass(frozen=True, slots=True)
class WorkbookEvidence:
    file_date: str
    row_count: int
    unique_npn_count: int
    duplicate_npn_values: int
    sha256: str
    byte_count: int
    content_type: str


@dataclass(frozen=True, slots=True)
class CanonicalGatewayRow:
    row_index: int
    producer: str
    producer_npn: str
    email: str | None
    upline: str | None
    upline_npn: str | None
    agency: str | None
    personal_ip_cents: int
    team_ip_cents: int
    first_leg_producer: str | None
    first_leg_npn: str | None
    first_leg_percent: float | None
    outside_first_leg_cents: int

    def as_payload(self) -> dict[str, Any]:
        return {
            "rowIndex": self.row_index,
            "producer": self.producer,
            "producerNpn": self.producer_npn,
            "email": self.email,
            "upline": self.upline,
            "uplineNpn": self.upline_npn,
            "agency": self.agency,
            "personalIpCents": self.personal_ip_cents,
            "teamIpCents": self.team_ip_cents,
            "firstLegProducer": self.first_leg_producer,
            "firstLegNpn": self.first_leg_npn,
            "firstLegPercent": self.first_leg_percent,
            "outsideFirstLegCents": self.outside_first_leg_cents,
        }


@dataclass(frozen=True, slots=True)
class GatewayArtifact:
    period_type: str
    source_filename: str
    content: bytes
    evidence: WorkbookEvidence
    coverage: dict[str, Any]
    generate_ms: int
    ready_ms: int
    download_ms: int
    total_ms: int
    rows: tuple[CanonicalGatewayRow, ...] = ()
    poll_count: int = 0


@dataclass(frozen=True, slots=True)
class ArtifactOutcome:
    period_type: str
    status: str
    artifact: GatewayArtifact | None = None
    error_code: str | None = None
    retryable: bool = False


@dataclass(frozen=True, slots=True)
class CurrentArtifactPull:
    status: str
    outcomes: dict[str, ArtifactOutcome]

    @property
    def artifacts(self) -> dict[str, GatewayArtifact]:
        return {
            period: outcome.artifact
            for period, outcome in self.outcomes.items()
            if outcome.artifact is not None
        }


@dataclass(frozen=True, slots=True)
class StatsBootstrap:
    endpoint: str
    uuid_token: str
    csrf_token: str
    spoofed_data: str
    report_uuid_hint: str


@dataclass(frozen=True, slots=True)
class WeeklyEntry:
    date: str
    week_ordinal: int
    month_id: str
    period_start: str
    period_end: str
    period_final: bool

    @property
    def gateway_label(self) -> str:
        month_name = self.month_id.split("-", 1)[0]
        return f"{month_name} Week {self.week_ordinal}"


@dataclass(frozen=True, slots=True)
class ReportIdentity:
    uuid: str
    gateway_date: str
    csrf_token: str
    current_week: WeeklyEntry
    available_week_count: int = 0
    oldest_week_date: str | None = None
    newest_week_date: str | None = None
    weekly_entries: tuple[WeeklyEntry, ...] = ()


@dataclass(frozen=True, slots=True)
class HistoricalPlanItem:
    period_type: str
    entry: WeeklyEntry


@dataclass(frozen=True, slots=True)
class HistoricalArtifactOutcome:
    period_type: str
    entry: WeeklyEntry
    status: str
    artifact: GatewayArtifact | None = None
    error_code: str | None = None
    retryable: bool = False


@dataclass(frozen=True, slots=True)
class HistoricalArtifactPull:
    status: str
    outcomes: tuple[HistoricalArtifactOutcome, ...]
    offset: int = 0
    total_count: int = 0


@dataclass(slots=True)
class ProviderMetrics:
    gateway_login_ms: int = 0
    gateway_handoff_ms: int = 0
    stats_autologin_ms: int = 0
    total_sync_ms: int = 0
    export_ms: dict[str, int] = field(default_factory=dict)


def _now_ms() -> int:
    return int(time.perf_counter() * 1000)


def _origin(url: str) -> str:
    parsed = urlparse(url)
    return f"{parsed.scheme}://{parsed.netloc}"


def _require_https(url: str, *, allowed_origin: str | None = None) -> str:
    raw = str(url or "").strip()
    if allowed_origin is not None:
        raw = urljoin(f"{allowed_origin.rstrip('/')}/", raw)
    parsed = urlparse(raw)
    if parsed.scheme != "https" or not parsed.netloc:
        raise GatewayContractError("Upstream returned a non-HTTPS URL.")
    if allowed_origin is not None and _origin(raw) != allowed_origin:
        raise GatewayContractError("Upstream returned an unexpected destination origin.")
    return raw


def _safe_json(response: Any, label: str) -> dict[str, Any]:
    try:
        payload = response.json()
    except Exception as exc:  # upstream boundary
        raise GatewayContractError(f"{label} did not return valid JSON.") from exc
    if not isinstance(payload, dict):
        raise GatewayContractError(f"{label} returned an unexpected JSON shape.")
    return payload


def classify_provider_error(error: BaseException) -> tuple[str, bool]:
    if isinstance(error, GatewayAuthenticationError):
        return "authentication_failed", False
    if isinstance(error, GatewayContractError):
        return "upstream_contract_changed", False
    if isinstance(error, GatewayTransientError):
        return "upstream_transient_failure", True
    if isinstance(error, (CurlTimeout, CurlConnectionError, TimeoutError, ConnectionError)):
        return "transport_failure", True
    return "unexpected_provider_failure", False


def _require_success_status(response: Any, label: str) -> None:
    status = int(response.status_code)
    if status in (401, 403):
        raise GatewayAuthenticationError(f"{label} was rejected by the upstream session.")
    if status in (408, 425, 429) or status >= 500:
        raise GatewayTransientError(f"{label} returned HTTP {status}.")
    if status < 200 or status >= 300:
        raise GatewayContractError(f"{label} returned unexpected HTTP {status}.")


def _remaining_timeout(deadline: float, cap_seconds: float, label: str) -> float:
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise GatewayTransientError(f"{label} timed out before completion.")
    return min(cap_seconds, max(0.05, remaining))


def _is_iso_date(value: str | None) -> bool:
    if not value or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return False
    try:
        return date.fromisoformat(value).isoformat() == value
    except ValueError:
        return False


def _month_bounds(month_id: str) -> tuple[str, str]:
    match = re.fullmatch(r"([A-Za-z]+)-(\d{4})", month_id or "")
    if not match:
        raise GatewayContractError("Gateway weekly navigation is missing its covered month identity.")
    names = (
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    )
    try:
        month = next(i for i, name in enumerate(names, start=1) if name.lower() == match.group(1).lower())
    except StopIteration as exc:
        raise GatewayContractError("Gateway weekly navigation returned an unknown month.") from exc
    year = int(match.group(2))
    start = date(year, month, 1)
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)
    end = date.fromordinal(next_month.toordinal() - 1)
    return start.isoformat(), end.isoformat()


class PageFactsParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.csrf: str | None = None
        self.report_uuid: str | None = None
        self.gateway_date: str | None = None
        self.has_login_email = False
        self.has_login_password = False
        self.has_otp = False
        self._stack: list[dict[str, str]] = []
        self._anchor: dict[str, Any] | None = None
        self.week_links: list[tuple[str, str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {str(key).lower(): str(value or "") for key, value in attrs}
        self._stack.append({"tag": tag.lower(), **values})
        if tag.lower() == "meta" and values.get("name", "").lower() == "csrf-token":
            self.csrf = values.get("content") or self.csrf
        if tag.lower() == "input":
            self.has_login_email |= values.get("id") == "login-email"
            self.has_login_password |= values.get("id") == "login-password"
            self.has_otp |= values.get("id") == "secretKey1"
        classes = set(values.get("class", "").split())
        if "downline-ip-report" in classes and self.report_uuid is None:
            self.report_uuid = values.get("data-uuid") or None
            self.gateway_date = values.get("data-week") or None
        if tag.lower() == "a" and "/ipreport/" in values.get("href", ""):
            month_id = ""
            for node in reversed(self._stack[:-1]):
                if re.fullmatch(r"[A-Za-z]+-\d{4}", node.get("id", "")):
                    month_id = node["id"]
                    break
            self._anchor = {"href": values.get("href", ""), "month_id": month_id, "text": []}

    def handle_data(self, data: str) -> None:
        if self._anchor is not None:
            self._anchor["text"].append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._anchor is not None:
            self.week_links.append(
                (
                    str(self._anchor["href"]),
                    " ".join(self._anchor["text"]).strip(),
                    str(self._anchor["month_id"]),
                )
            )
            self._anchor = None
        for index in range(len(self._stack) - 1, -1, -1):
            if self._stack[index].get("tag") == tag.lower():
                del self._stack[index:]
                break


def _page_facts(html: str) -> PageFactsParser:
    parser = PageFactsParser()
    parser.feed(html)
    return parser


def _decode_js_double_quoted(raw: str) -> str:
    try:
        return json.loads(f'"{raw}"')
    except json.JSONDecodeError as exc:
        raise GatewayContractError("StatS handoff spoofed_data string could not be decoded.") from exc


def parse_stats_bootstrap(html: str, stats_origin: str) -> StatsBootstrap:
    request_match = re.search(
        r"url\s*:\s*([\"'])([^\"']*/api/auth/user/auto-login/([0-9a-f-]{36}))\1"
        r"[\s\S]{0,2000}?data\s*:\s*\{([\s\S]{0,1800}?)\}",
        html,
        re.I,
    )
    spoof_match = re.search(r'var\s+spoofed_data\s*=\s*"((?:\\.|[^"\\])*)"', html, re.I)
    if not request_match or not spoof_match:
        raise GatewayContractError("StatS auto-login bootstrap contract changed or is incomplete.")
    data_block = request_match.group(4)
    uuid_match = re.search(r"[\"']?uuid[\"']?\s*:\s*([\"'])([^\"']+)\1", data_block, re.I)
    token_match = re.search(r"_token\s*:\s*([\"'])([^\"']+)\1", data_block, re.I)
    if not uuid_match or not token_match:
        raise GatewayContractError("StatS auto-login payload fields changed or are missing.")
    endpoint = _require_https(urljoin(stats_origin, request_match.group(2)), allowed_origin=stats_origin)
    return StatsBootstrap(
        endpoint=endpoint,
        uuid_token=uuid_match.group(2),
        csrf_token=token_match.group(2),
        spoofed_data=_decode_js_double_quoted(spoof_match.group(1)),
        report_uuid_hint=request_match.group(3),
    )


def _weekly_entries(facts: PageFactsParser, uuid: str) -> list[WeeklyEntry]:
    raw: list[tuple[str, int, str, str, str]] = []
    for href, text, month_id in facts.week_links:
        date_match = re.search(rf"/ipreport/{re.escape(uuid)}/(\d{{4}}-\d{{2}}-\d{{2}})(?:[?#].*)?$", href)
        week_match = re.search(r"Week\s+(\d+)", text, re.I)
        if not date_match or not week_match or not month_id:
            continue
        gateway_date = date_match.group(1)
        ordinal = int(week_match.group(1))
        if not _is_iso_date(gateway_date) or ordinal < 1 or ordinal > 8:
            continue
        period_start, period_end = _month_bounds(month_id)
        if gateway_date < period_start or gateway_date > period_end:
            raise GatewayContractError("Gateway weekly navigation date falls outside its covered month.")
        row = (gateway_date, ordinal, month_id, period_start, period_end)
        if row not in raw:
            raw.append(row)
    raw.sort(key=lambda row: (row[0], row[1], row[3]))
    max_by_month: dict[str, int] = {}
    for _, ordinal, _, period_start, _ in raw:
        max_by_month[period_start] = max(max_by_month.get(period_start, 0), ordinal)
    return [
        WeeklyEntry(
            date=gateway_date,
            week_ordinal=ordinal,
            month_id=month_id,
            period_start=period_start,
            period_end=period_end,
            period_final=ordinal == max_by_month[period_start],
        )
        for gateway_date, ordinal, month_id, period_start, period_end in raw
    ]


def parse_report_identity(html: str, expected_uuid: str) -> ReportIdentity:
    facts = _page_facts(html)
    if not facts.report_uuid or not _is_iso_date(facts.gateway_date) or not facts.csrf:
        raise GatewayContractError("StatS current report page did not expose the expected report identity.")
    if facts.report_uuid != expected_uuid:
        raise GatewayContractError("StatS report identity did not match the auto-login account context.")
    entries = _weekly_entries(facts, expected_uuid)
    current = next((entry for entry in entries if entry.date == facts.gateway_date), None)
    if current is None:
        raise GatewayContractError("Gateway current weekly identity was not found in report navigation.")
    return ReportIdentity(
        uuid=facts.report_uuid,
        gateway_date=facts.gateway_date or "",
        csrf_token=facts.csrf,
        current_week=current,
        available_week_count=len(entries),
        oldest_week_date=min((entry.date for entry in entries), default=None),
        newest_week_date=max((entry.date for entry in entries), default=None),
        weekly_entries=tuple(entries),
    )


def historical_plan(
    entries: tuple[WeeklyEntry, ...] | list[WeeklyEntry],
    from_date: str,
    through_date: str,
) -> tuple[HistoricalPlanItem, ...]:
    if not _is_iso_date(from_date) or not _is_iso_date(through_date) or through_date < from_date:
        raise ValueError("Gateway historical pull requires a valid YYYY-MM-DD range.")
    eligible = [entry for entry in entries if from_date <= entry.date <= through_date]
    if not eligible:
        return ()
    latest = eligible[-1]
    plan: list[HistoricalPlanItem] = []
    for entry in eligible:
        plan.append(HistoricalPlanItem(period_type="wtd", entry=entry))
        if entry.period_final:
            plan.append(HistoricalPlanItem(period_type="mtd", entry=entry))
            plan.append(HistoricalPlanItem(period_type="ytd", entry=entry))
    if not latest.period_final:
        plan.append(HistoricalPlanItem(period_type="mtd", entry=latest))
        plan.append(HistoricalPlanItem(period_type="ytd", entry=latest))
    return tuple(plan)


def _cookie_domain_matches_host(domain: str, host: str) -> bool:
    normalized = domain.strip().lower().lstrip(".").rstrip(".")
    target = host.strip().lower().rstrip(".")
    if not normalized or normalized.count(".") < 1:
        return False
    return target == normalized or target.endswith(f".{normalized}")


def _validated_session_cookies(
    session_state: Any,
    *,
    trusted_hosts: tuple[str, ...],
    now: float | None = None,
    require_nonempty: bool = True,
) -> tuple[dict[str, Any], ...]:
    if not isinstance(session_state, dict) or not isinstance(session_state.get("cookies"), list):
        raise GatewayContractError("Gateway HTTP session state is malformed.")
    raw_cookies = session_state["cookies"]
    if len(raw_cookies) > 512:
        raise GatewayContractError("Gateway HTTP session state contains too many cookies.")
    current_time = time.time() if now is None else now
    accepted: list[dict[str, Any]] = []
    for raw in raw_cookies:
        if not isinstance(raw, dict):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie.")
        name = raw.get("name")
        value = raw.get("value")
        domain = raw.get("domain")
        path = raw.get("path", "/")
        if not isinstance(name, str) or not name or len(name) > 256 or re.search(r"[\x00-\x20\x7f]", name):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie name.")
        if not isinstance(value, str) or len(value) > 16384 or re.search(r"[\x00\r\n]", value):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie value.")
        if (
            not isinstance(domain, str)
            or len(domain) > 255
            or domain != domain.strip()
            or re.search(r"[\x00-\x20\x7f]", domain)
        ):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie domain.")
        if not any(_cookie_domain_matches_host(domain, host) for host in trusted_hosts):
            raise GatewayContractError("Gateway HTTP session state contains a cookie outside trusted Gateway origins.")
        if not isinstance(path, str) or not path.startswith("/") or len(path) > 2048 or re.search(r"[\x00\r\n]", path):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie path.")
        secure = raw.get("secure", False)
        http_only = raw.get("httpOnly", False)
        if not isinstance(secure, bool) or not isinstance(http_only, bool):
            raise GatewayContractError("Gateway HTTP session state contains malformed cookie flags.")
        same_site = raw.get("sameSite")
        if same_site is not None and same_site not in {"Strict", "Lax", "None"}:
            raise GatewayContractError("Gateway HTTP session state contains an unsupported cookie SameSite value.")
        expires = raw.get("expires", -1)
        if isinstance(expires, bool) or not isinstance(expires, (int, float)) or not math.isfinite(float(expires)):
            raise GatewayContractError("Gateway HTTP session state contains a malformed cookie expiry.")
        if float(expires) >= 0 and float(expires) <= current_time:
            continue
        accepted.append({
            "name": name,
            "value": value,
            "domain": domain.lower(),
            "path": path,
            "secure": secure,
            "httpOnly": http_only,
            "sameSite": same_site,
            "expires": float(expires) if float(expires) >= 0 else -1,
        })
    if require_nonempty and not accepted:
        raise GatewayAuthenticationError("Gateway HTTP session has no usable cookies and requires reconnect.")
    return tuple(accepted)


def _validated_auth_challenge_state(state: Any) -> dict[str, Any]:
    if not isinstance(state, dict) or state.get("version") != 1:
        raise GatewayContractError("Gateway two-factor challenge state is malformed.")
    email = state.get("email")
    token = state.get("token")
    phase = state.get("phase")
    kind = state.get("type")
    verified = state.get("verified")
    value = state.get("value")
    session_state = state.get("sessionState")
    if not isinstance(email, str) or not email.strip() or len(email) > 320 or re.search(r"[\x00\r\n]", email):
        raise GatewayContractError("Gateway two-factor challenge state has an invalid email.")
    if not isinstance(token, str) or not token or len(token) > 8192 or re.search(r"[\x00\r\n]", token):
        raise GatewayContractError("Gateway two-factor challenge state has an invalid token.")
    if phase not in {"otp_required", "setup_phone", "setup_email"}:
        raise GatewayContractError("Gateway two-factor challenge state has an invalid phase.")
    if kind not in {"all", "phone", "email"}:
        raise GatewayContractError("Gateway two-factor challenge state has an invalid type.")
    if isinstance(verified, bool) or verified not in {0, 1}:
        raise GatewayContractError("Gateway two-factor challenge state has an invalid verification flag.")
    if value is not None and (
        not isinstance(value, str) or not value.strip() or len(value) > 2048 or re.search(r"[\x00\r\n]", value)
    ):
        raise GatewayContractError("Gateway two-factor challenge state has an invalid setup value.")
    if not isinstance(session_state, dict):
        raise GatewayContractError("Gateway two-factor challenge state is missing its HTTP session.")
    if (phase, kind, verified) not in {
        ("setup_phone", "phone", 0),
        ("setup_email", "email", 1),
        ("otp_required", "phone", 0),
        ("otp_required", "email", 1),
        ("otp_required", "all", 1),
    }:
        raise GatewayContractError("Gateway two-factor challenge state has an inconsistent phase.")
    return {
        "version": 1,
        "email": email.strip(),
        "token": token,
        "phase": phase,
        "type": kind,
        "verified": verified,
        "value": value.strip() if isinstance(value, str) else None,
        "sessionState": session_state,
    }


def artifact_csv_text(artifact: GatewayArtifact) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(REQUIRED_COLUMNS)
    file_date = date.fromisoformat(artifact.evidence.file_date)
    file_date_text = f"{file_date.month}/{file_date.day}/{file_date.year}"

    def money(cents: int) -> str:
        return f"${Decimal(cents) / Decimal(100):,.2f}"

    for row in artifact.rows:
        writer.writerow([
            file_date_text,
            row.producer,
            row.producer_npn,
            row.email or "",
            row.upline or "",
            row.upline_npn or "",
            row.agency or "",
            money(row.personal_ip_cents),
            money(row.team_ip_cents),
            row.first_leg_producer or "",
            row.first_leg_npn or "",
            "" if row.first_leg_percent is None else f"{row.first_leg_percent:.2f}%",
            money(row.outside_first_leg_cents),
        ])
    return output.getvalue().rstrip("\n")


def _normalize_npn(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\D", "", str(value))


def _normalize_file_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date().isoformat()
            if isinstance(converted, date):
                return converted.isoformat()
        except Exception:
            return None
    text = str(value or "").strip()
    if _is_iso_date(text):
        return text
    for fmt in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _clean_text(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _money_to_cents(value: Any, *, label: str) -> int:
    if value is None or value == "":
        return 0
    text = str(value).replace("$", "").replace(",", "").strip()
    if not text:
        return 0
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise GatewayContractError(f"{label} contains an invalid money value.") from exc
    if not amount.is_finite() or amount < 0:
        raise GatewayContractError(f"{label} contains an invalid money value.")
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _percent_value(value: Any, *, label: str) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise GatewayContractError(f"{label} contains an invalid percent value.")
    is_percent_text = isinstance(value, str) and "%" in value
    text = str(value).replace("%", "").strip()
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise GatewayContractError(f"{label} contains an invalid percent value.") from exc
    if not number.is_finite():
        raise GatewayContractError(f"{label} contains an invalid percent value.")
    if not is_percent_text and isinstance(value, (int, float, Decimal)) and Decimal("0") <= number <= Decimal("1"):
        number *= 100
    if number < 0 or number > 100:
        raise GatewayContractError(f"{label} contains a percent outside 0-100.")
    return float(number)


def _parse_workbook(
    content: bytes,
    *,
    period: str,
    identity: ReportIdentity,
    content_type: str,
) -> tuple[WorkbookEvidence, tuple[CanonicalGatewayRow, ...]]:
    if period not in PERIOD_CODES:
        raise ValueError(f"Unsupported Gateway period: {period}")
    if len(content) < 4 or content[:4] != b"PK\x03\x04":
        raise GatewayContractError(f"{period.upper()} response is not an XLSX/ZIP workbook.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise GatewayContractError(f"{period.upper()} workbook could not be opened.") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        source_rows = sheet.iter_rows(values_only=True)
        headers_raw = next(source_rows, None)
        if headers_raw is None:
            raise GatewayContractError(f"{period.upper()} workbook is empty.")
        headers = [str(value or "").lstrip("\ufeff").strip() for value in headers_raw]
        index = {name: position for position, name in enumerate(headers)}
        missing = [name for name in REQUIRED_COLUMNS if name not in index]
        if missing:
            raise GatewayContractError(
                f"{period.upper()} workbook is missing required columns: {', '.join(missing)}"
            )

        file_dates: set[str] = set()
        npns: list[str] = []
        parsed_rows: list[CanonicalGatewayRow] = []
        for source_row_index, row in enumerate(source_rows, start=2):
            producer = _clean_text(row[index["Producer"]])
            npn = _normalize_npn(row[index["Producer NPN"]])
            if not producer or not npn:
                continue
            file_date = _normalize_file_date(row[index["File Date"]])
            if not file_date:
                raise GatewayContractError(f"{period.upper()} workbook contains an invalid File Date.")
            file_dates.add(file_date)
            npns.append(npn)
            parsed_rows.append(CanonicalGatewayRow(
                row_index=source_row_index,
                producer=producer,
                producer_npn=npn,
                email=_clean_text(row[index["Email"]]),
                upline=_clean_text(row[index["Upline"]]),
                upline_npn=_clean_text(_normalize_npn(row[index["Upline NPN"]])),
                agency=_clean_text(row[index["Agency"]]),
                personal_ip_cents=_money_to_cents(row[index["All Lines Personal"]], label=f"{period.upper()} row {source_row_index}"),
                team_ip_cents=_money_to_cents(row[index["All Lines Team"]], label=f"{period.upper()} row {source_row_index}"),
                first_leg_producer=_clean_text(row[index["1st Leg Producer"]]),
                first_leg_npn=_clean_text(_normalize_npn(row[index["1st Leg NPN"]])),
                first_leg_percent=_percent_value(row[index["1st Leg %"]], label=f"{period.upper()} row {source_row_index}"),
                outside_first_leg_cents=_money_to_cents(row[index["Outside 1st Leg"]], label=f"{period.upper()} row {source_row_index}"),
            ))
        if not parsed_rows or not file_dates:
            raise GatewayContractError(f"{period.upper()} workbook contains no usable producer rows.")
        if len(file_dates) != 1:
            raise GatewayContractError(f"{period.upper()} workbook contains inconsistent File Date values.")
        file_date = next(iter(file_dates))
        if period == "wtd":
            start = date.fromisoformat(file_date)
            end = date.fromisoformat(identity.gateway_date)
            if (end - start).days not in range(0, 7):
                raise GatewayContractError("WTD File Date is inconsistent with the weekly report identity.")
        elif period == "mtd":
            if file_date != identity.current_week.period_start:
                raise GatewayContractError("MTD File Date does not match the covered month start.")
        else:
            expected = f"{identity.gateway_date[:4]}-01-01"
            if file_date != expected:
                raise GatewayContractError("YTD File Date does not match the covered year start.")
        counts = Counter(npns)
        evidence = WorkbookEvidence(
            file_date=file_date,
            row_count=len(parsed_rows),
            unique_npn_count=len(counts),
            duplicate_npn_values=sum(1 for count in counts.values() if count > 1),
            sha256=hashlib.sha256(content).hexdigest(),
            byte_count=len(content),
            content_type=content_type,
        )
        return evidence, tuple(parsed_rows)
    finally:
        workbook.close()


def validate_workbook(
    content: bytes,
    *,
    period: str,
    identity: ReportIdentity,
    content_type: str,
) -> WorkbookEvidence:
    evidence, _ = _parse_workbook(
        content,
        period=period,
        identity=identity,
        content_type=content_type,
    )
    return evidence


class GatewayHttpProvider:
    """HTTP-first adapter for the observed private Gateway/StatS web contract.

    Raw endpoint details, cookies, tokens and handoff parsing intentionally stay
    inside this class. Callers receive only validated XLSX artifacts + evidence.
    """

    def __init__(self, config: HttpGatewayConfig | None = None) -> None:
        self.config = config or HttpGatewayConfig()
        self.session = AsyncSession(impersonate=self.config.impersonate, max_clients=8)
        self.gateway_access_token: str | None = None
        self.stats_access_token: str | None = None
        self.identity: ReportIdentity | None = None
        self.report_url: str | None = None
        self.metrics = ProviderMetrics()

    async def close(self) -> None:
        await self.session.close()

    async def __aenter__(self) -> GatewayHttpProvider:
        return self

    async def __aexit__(self, *_: object) -> None:
        await self.close()

    async def bootstrap_probe(self) -> dict[str, Any]:
        started = _now_ms()
        response = await self.session.get(
            f"{self.config.gateway_origin}/login",
            timeout=self.config.request_timeout_seconds,
        )
        _require_success_status(response, "Gateway login bootstrap")
        facts = _page_facts(response.text)
        if not (facts.has_login_email and facts.has_login_password):
            raise GatewayContractError("Gateway login bootstrap did not expose the expected login form.")
        return {
            "ok": True,
            "transport": "curl_cffi/chrome",
            "status": response.status_code,
            "latency_ms": _now_ms() - started,
        }

    async def contract_probe(self, *, email: str, password: str) -> dict[str, Any]:
        """Exercise the scheduled no-export compatibility surface.

        This intentionally stops after login, Gateway->StatS handoff, StatS
        session establishment and current report identity discovery. It never
        starts WTD/MTD/YTD generation.
        """
        started = _now_ms()
        await self.bootstrap_probe()
        await self._gateway_login(email=email.strip(), password=password)
        self.identity = await self._gateway_to_stats_handoff()
        return {
            "ok": True,
            "transport": "curl_cffi/chrome",
            "report_identity": True,
            "historical_report_identity": self.identity.available_week_count > 1 if self.identity else False,
            "navigation_week_count": self.identity.available_week_count if self.identity else 0,
            "total_ms": _now_ms() - started,
            "gateway_login_ms": self.metrics.gateway_login_ms,
            "gateway_handoff_ms": self.metrics.gateway_handoff_ms,
            "stats_autologin_ms": self.metrics.stats_autologin_ms,
        }

    async def contract_probe_storage_state(self, storage_state: Any) -> dict[str, Any]:
        started = _now_ms()
        await self.establish_session_from_state(storage_state)
        return {
            "ok": True,
            "transport": "curl_cffi/chrome",
            "report_identity": True,
            "historical_report_identity": self.identity.available_week_count > 1 if self.identity else False,
            "navigation_week_count": self.identity.available_week_count if self.identity else 0,
            "total_ms": _now_ms() - started,
            "gateway_login_ms": self.metrics.gateway_login_ms,
            "gateway_handoff_ms": self.metrics.gateway_handoff_ms,
            "stats_autologin_ms": self.metrics.stats_autologin_ms,
        }

    async def establish_session(self, *, email: str, password: str) -> ReportIdentity:
        if not email.strip() or not password:
            raise GatewayAuthenticationError("Gateway email and password are required.")
        await self.contract_probe(email=email, password=password)
        if self.identity is None:
            raise GatewayContractError("Gateway contract probe completed without report identity.")
        return self.identity

    def _trusted_session_hosts(self) -> tuple[str, str]:
        gateway_host = urlparse(self.config.gateway_origin).hostname
        stats_host = urlparse(self.config.stats_origin).hostname
        if not gateway_host or not stats_host:
            raise GatewayContractError("Trusted Gateway origins are missing hostnames.")
        return gateway_host, stats_host

    def export_session_state(self, *, require_nonempty: bool = True) -> dict[str, Any]:
        cookies: list[dict[str, Any]] = []
        for cookie in self.session.cookies.jar:
            rest = getattr(cookie, "_rest", {}) or {}
            same_site = rest.get("SameSite") or rest.get("samesite")
            if same_site not in {"Strict", "Lax", "None"}:
                same_site = None
            cookies.append({
                "name": str(cookie.name),
                "value": str(cookie.value),
                "domain": str(cookie.domain),
                "path": str(cookie.path or "/"),
                "expires": int(cookie.expires) if cookie.expires is not None else -1,
                "httpOnly": bool(rest.get("HttpOnly") is not None or rest.get("httponly") is not None),
                "secure": bool(cookie.secure),
                **({"sameSite": same_site} if same_site is not None else {}),
            })
        validated = _validated_session_cookies(
            {"cookies": cookies},
            trusted_hosts=self._trusted_session_hosts(),
            require_nonempty=require_nonempty,
        )
        return {"cookies": [dict(cookie) for cookie in validated]}

    def import_session_state(self, session_state: Any, *, require_nonempty: bool = True) -> int:
        if (
            isinstance(session_state, dict)
            and session_state.get("version") == 2
            and isinstance(session_state.get("sessionState"), dict)
        ):
            session_state = session_state["sessionState"]
        cookies = _validated_session_cookies(
            session_state,
            trusted_hosts=self._trusted_session_hosts(),
            require_nonempty=require_nonempty,
        )
        self.session.cookies.clear()
        for cookie in cookies:
            self.session.cookies.set(
                cookie["name"],
                cookie["value"],
                domain=cookie["domain"],
                path=cookie["path"],
                secure=cookie["secure"],
            )
        return len(cookies)

    # Backward-compatible aliases for runner sessions originally captured as
    # Playwright storageState. The HTTP adapter only consumes their cookie list.
    def import_storage_state(self, storage_state: Any) -> int:
        return self.import_session_state(storage_state)

    async def establish_session_from_state(self, session_state: Any) -> ReportIdentity:
        self.gateway_access_token = None
        self.stats_access_token = None
        self.identity = None
        self.report_url = None
        self.import_session_state(session_state)
        started = _now_ms()
        dashboard = await self.session.get(
            f"{self.config.gateway_origin}/dashboard",
            timeout=self.config.request_timeout_seconds,
        )
        self.metrics.gateway_login_ms = _now_ms() - started
        _require_success_status(dashboard, "Gateway dashboard session resume")
        facts = _page_facts(dashboard.text)
        if facts.has_login_email or facts.has_login_password or facts.has_otp:
            raise GatewayAuthenticationError("Gateway HTTP session expired and requires reconnect.")
        if not facts.csrf:
            raise GatewayContractError("Gateway dashboard session resume is missing the expected CSRF token.")
        self.identity = await self._gateway_to_stats_handoff(dashboard=dashboard)
        return self.identity

    async def establish_session_from_storage_state(self, storage_state: Any) -> ReportIdentity:
        return await self.establish_session_from_state(storage_state)

    async def _gateway_login_response(self, *, email: str, password: str) -> dict[str, Any]:
        started = _now_ms()
        response = await self.session.post(
            f"{self.config.gateway_origin}/api/auth/user/login",
            data={"email": email, "password": password},
            headers={"X-Requested-With": "XMLHttpRequest"},
            timeout=self.config.request_timeout_seconds,
        )
        self.metrics.gateway_login_ms = _now_ms() - started
        if response.status_code in (400, 401, 403, 422):
            raise GatewayAuthenticationError("Gateway login was rejected.")
        if response.status_code == 429 or response.status_code >= 500:
            raise GatewayTransientError(f"Gateway login returned HTTP {response.status_code}.")
        payload = _safe_json(response, "Gateway login")
        if not payload.get("status"):
            raise GatewayAuthenticationError("Gateway login was rejected.")
        return payload

    async def _gateway_auth_user(self, access_token: str) -> None:
        if not isinstance(access_token, str) or not access_token or len(access_token) > 16384:
            raise GatewayContractError("Gateway authentication completed without a usable access token.")
        self.gateway_access_token = access_token
        auth_response = await self.session.get(
            f"{self.config.gateway_origin}/api/auth/user/auth-user",
            headers={"Authorization": f"Bearer {access_token}"},
            timeout=self.config.request_timeout_seconds,
        )
        _require_success_status(auth_response, "Gateway authenticated identity probe")
        auth_payload = _safe_json(auth_response, "Gateway auth-user")
        if not auth_payload.get("status"):
            raise GatewayAuthenticationError("Gateway authenticated identity probe failed.")

    async def _finish_http_connect(self, access_token: str) -> dict[str, Any]:
        await self._gateway_auth_user(access_token)
        self.identity = await self._gateway_to_stats_handoff()
        return {
            "status": "connected",
            "session": {
                "version": 2,
                "sessionState": self.export_session_state(),
            },
            "evidence": {
                "navigationWeekCount": self.identity.available_week_count,
                "historicalNavigation": self.identity.available_week_count > 1,
            },
        }

    def _challenge_result(
        self,
        *,
        email: str,
        token: str,
        phase: str,
        kind: str,
        verified: int,
        value: str | None = None,
    ) -> dict[str, Any]:
        return {
            "status": phase,
            "state": {
                "version": 1,
                "email": email,
                "token": token,
                "phase": phase,
                "type": kind,
                "verified": verified,
                **({"value": value} if value is not None else {}),
                "sessionState": self.export_session_state(require_nonempty=False),
            },
        }

    def _rejected_challenge_result(
        self,
        *,
        email: str,
        token: str,
        phase: str,
        kind: str,
        verified: int,
        value: str | None = None,
    ) -> dict[str, Any]:
        result = self._challenge_result(
            email=email,
            token=token,
            phase=phase,
            kind=kind,
            verified=verified,
            value=value,
        )
        result["rejected"] = True
        result["errorCode"] = "authentication_failed"
        return result

    async def connect_begin(self, *, email: str, password: str) -> dict[str, Any]:
        email = email.strip()
        if not email or len(email) > 320 or re.search(r"[\x00\r\n]", email) or not password or len(password) > 4096:
            raise GatewayAuthenticationError("Gateway email and password are required.")
        await self.bootstrap_probe()
        payload = await self._gateway_login_response(email=email, password=password)
        if not payload.get("twoFaEnabled"):
            token = payload.get("access_token")
            if not isinstance(token, str) or not token:
                raise GatewayContractError("Gateway login succeeded without the expected access token.")
            return await self._finish_http_connect(token)
        remember_token = payload.get("remember_token")
        if not isinstance(remember_token, str) or not remember_token or len(remember_token) > 8192:
            raise GatewayContractError("Gateway two-factor login did not return its challenge token.")
        if payload.get("showSetupScreen"):
            return self._challenge_result(
                email=email,
                token=remember_token,
                phase="setup_phone",
                kind="phone",
                verified=0,
            )
        if payload.get("showVerifyScreen"):
            return self._challenge_result(
                email=email,
                token=remember_token,
                phase="otp_required",
                kind="all",
                verified=1,
            )
        raise GatewayContractError("Gateway two-factor login returned an unsupported challenge state.")

    async def _two_factor_post(self, path: str, data: dict[str, Any], label: str) -> dict[str, Any]:
        response = await self.session.post(
            f"{self.config.gateway_origin}{path}",
            data=data,
            headers={"X-Requested-With": "XMLHttpRequest"},
            timeout=self.config.request_timeout_seconds,
        )
        if response.status_code in (400, 401, 403, 409, 422):
            raise GatewayAuthenticationError(f"Gateway {label} was rejected.")
        if response.status_code in (408, 425, 429) or response.status_code >= 500:
            raise GatewayTransientError(f"Gateway {label} returned HTTP {response.status_code}.")
        if response.status_code < 200 or response.status_code >= 300:
            raise GatewayContractError(f"Gateway {label} returned unexpected HTTP {response.status_code}.")
        return _safe_json(response, f"Gateway {label}")

    async def connect_step(
        self,
        *,
        state: Any,
        action: str | None = None,
        value: str | None = None,
        code: str | None = None,
    ) -> dict[str, Any]:
        challenge = _validated_auth_challenge_state(state)
        self.import_session_state(challenge["sessionState"], require_nonempty=False)
        email = challenge["email"]
        token = challenge["token"]
        phase = challenge["phase"]
        kind = challenge["type"]
        verified = challenge["verified"]
        stored_value = challenge["value"]
        normalized_action = (action or "").strip().lower()

        if phase in {"setup_phone", "setup_email"}:
            if normalized_action not in {"", "register", "submit", "continue"}:
                raise GatewayContractError("Gateway two-factor setup action is invalid.")
            setup_value = value.strip() if isinstance(value, str) else ""
            if not setup_value or len(setup_value) > 2048 or re.search(r"[\x00\r\n]", setup_value):
                raise GatewayAuthenticationError("Gateway two-factor setup value is required.")
            try:
                payload = await self._two_factor_post(
                    "/api/two-fa/register",
                    {
                        "email": email,
                        "userType": "user",
                        "token": token,
                        "value": setup_value,
                        "type": kind,
                    },
                    "two-factor registration",
                )
            except GatewayAuthenticationError:
                return self._rejected_challenge_result(
                    email=email,
                    token=token,
                    phase=phase,
                    kind=kind,
                    verified=verified,
                )
            if payload.get("status") is False:
                return self._rejected_challenge_result(
                    email=email,
                    token=token,
                    phase=phase,
                    kind=kind,
                    verified=verified,
                )
            return self._challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=setup_value,
            )

        if phase != "otp_required":
            raise GatewayContractError("Gateway two-factor challenge phase is unsupported.")

        if normalized_action == "resend":
            data: dict[str, Any] = {
                "email": email,
                "userType": "user",
                "token": token,
                "type": kind,
            }
            if stored_value is not None:
                data["value"] = stored_value
            try:
                payload = await self._two_factor_post("/api/two-fa/register", data, "two-factor resend")
            except GatewayAuthenticationError:
                return self._rejected_challenge_result(
                    email=email,
                    token=token,
                    phase="otp_required",
                    kind=kind,
                    verified=verified,
                    value=stored_value,
                )
            if payload.get("status") is False:
                return self._rejected_challenge_result(
                    email=email,
                    token=token,
                    phase="otp_required",
                    kind=kind,
                    verified=verified,
                    value=stored_value,
                )
            return self._challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )

        if normalized_action not in {"", "verify", "submit", "continue"}:
            raise GatewayContractError("Gateway two-factor verification action is invalid.")
        otp = code.strip() if isinstance(code, str) else ""
        if not re.fullmatch(r"\d{6}", otp):
            return self._rejected_challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )
        verify_data: dict[str, Any] = {
            "email": email,
            "userType": "user",
            "token": token,
            "secretKey": otp,
            "verified": verified,
            "type": kind,
        }
        if stored_value is not None:
            verify_data["value"] = stored_value
        try:
            verify_payload = await self._two_factor_post(
                "/api/two-fa/verify",
                verify_data,
                "two-factor verification",
            )
        except GatewayAuthenticationError:
            return self._rejected_challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )
        if not verify_payload.get("status"):
            return self._rejected_challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )

        if verified == 0:
            # This mirrors the observed Gateway client: successful phone setup
            # always advances to email setup before handleLogin is allowed.
            return self._challenge_result(
                email=email,
                token=token,
                phase="setup_email",
                kind="email",
                verified=1,
            )

        try:
            login_payload = await self._two_factor_post(
                "/api/two-fa/handleLogin",
                {"email": email, "userType": "user", "token": token},
                "two-factor login completion",
            )
        except GatewayAuthenticationError:
            return self._rejected_challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )
        access_token = login_payload.get("access_token")
        if not login_payload.get("status") or not login_payload.get("active"):
            return self._rejected_challenge_result(
                email=email,
                token=token,
                phase="otp_required",
                kind=kind,
                verified=verified,
                value=stored_value,
            )
        if not isinstance(access_token, str) or not access_token:
            raise GatewayContractError("Gateway two-factor login completion omitted its access token.")
        return await self._finish_http_connect(access_token)

    async def _gateway_login(self, *, email: str, password: str) -> None:
        payload = await self._gateway_login_response(email=email, password=password)
        if payload.get("twoFaEnabled"):
            raise GatewayAuthenticationError("Gateway contract probe requires interactive two-factor completion.")
        token = payload.get("access_token")
        if not isinstance(token, str) or not token:
            raise GatewayContractError("Gateway login succeeded without the expected access token.")
        await self._gateway_auth_user(token)

    async def _gateway_to_stats_handoff(self, *, dashboard: Any | None = None) -> ReportIdentity:
        if dashboard is None:
            dashboard = await self.session.get(
                f"{self.config.gateway_origin}/dashboard",
                timeout=self.config.request_timeout_seconds,
            )
            _require_success_status(dashboard, "Gateway dashboard bootstrap")
        dashboard_facts = _page_facts(dashboard.text)
        if dashboard_facts.has_login_email or dashboard_facts.has_login_password or dashboard_facts.has_otp:
            raise GatewayAuthenticationError("Gateway dashboard session was not established after login.")
        if not dashboard_facts.csrf:
            raise GatewayContractError("Gateway dashboard is missing the expected CSRF token.")
        started = _now_ms()
        handoff_headers = {
            "X-CSRF-TOKEN": dashboard_facts.csrf,
            "X-Requested-With": "XMLHttpRequest",
        }
        if self.gateway_access_token:
            handoff_headers["Authorization"] = f"Bearer {self.gateway_access_token}"
        handoff = await self.session.post(
            f"{self.config.gateway_origin}/api/autologinUser/ipreport",
            data={"base_url": self.config.stats_origin},
            headers=handoff_headers,
            timeout=self.config.request_timeout_seconds,
        )
        self.metrics.gateway_handoff_ms = _now_ms() - started
        _require_success_status(handoff, "Gateway IP Report handoff")
        payload = _safe_json(handoff, "Gateway IP Report handoff")
        signed_url = payload.get("url")
        if not payload.get("status") or not isinstance(signed_url, str):
            raise GatewayContractError("Gateway IP Report handoff did not return a usable destination.")
        signed_url = _require_https(signed_url, allowed_origin=self.config.stats_origin)
        started = _now_ms()
        bootstrap_response = await self.session.get(signed_url, timeout=self.config.request_timeout_seconds)
        _require_success_status(bootstrap_response, "StatS handoff bootstrap")
        bootstrap = parse_stats_bootstrap(bootstrap_response.text, self.config.stats_origin)
        auto_login = await self.session.post(
            bootstrap.endpoint,
            data={
                "uuid": bootstrap.uuid_token,
                "_token": bootstrap.csrf_token,
                "spoofed_data": bootstrap.spoofed_data,
            },
            headers={"X-Requested-With": "XMLHttpRequest", "Referer": signed_url},
            timeout=self.config.request_timeout_seconds,
        )
        self.metrics.stats_autologin_ms = _now_ms() - started
        _require_success_status(auto_login, "StatS auto-login")
        payload = _safe_json(auto_login, "StatS auto-login")
        token = payload.get("access_token")
        if not payload.get("status") or not isinstance(token, str) or not token:
            raise GatewayAuthenticationError("StatS auto-login failed.")
        self.stats_access_token = token
        auth_response = await self.session.get(
            f"{self.config.stats_origin}/api/auth/user/auth-user",
            headers={"Authorization": f"Bearer {token}"},
            timeout=self.config.request_timeout_seconds,
        )
        _require_success_status(auth_response, "StatS authenticated identity probe")
        auth_payload = _safe_json(auth_response, "StatS auth-user")
        if not auth_payload.get("status"):
            raise GatewayAuthenticationError("StatS authenticated identity probe failed.")
        report_url = f"{self.config.stats_origin}/ipreport/{bootstrap.report_uuid_hint}"
        report_response = await self.session.get(report_url, timeout=self.config.request_timeout_seconds)
        _require_success_status(report_response, "StatS report identity")
        identity = parse_report_identity(report_response.text, bootstrap.report_uuid_hint)
        self.report_url = report_url
        return identity

    def _export_params(self, period: str, *, identity: ReportIdentity | None = None) -> dict[str, str]:
        active_identity = identity or self.identity
        if period not in PERIOD_CODES or not active_identity:
            raise GatewayContractError("StatS report identity is not ready.")
        encoded_uuid = base64.b64encode(
            json.dumps({"uuid": active_identity.uuid}, separators=(",", ":")).encode("utf-8")
        ).decode("ascii")
        return {
            "selected_period": PERIOD_CODES[period],
            "selected_producer": "",
            "selected_agency": "",
            "selected_type": "",
            "selected_date": active_identity.gateway_date,
            "uuid": encoded_uuid,
            "_token": active_identity.csrf_token,
        }

    async def _fetch_workbook(
        self,
        content_url: str,
        *,
        period: str,
        report_date: str | None = None,
        deadline: float | None = None,
    ) -> tuple[bytes, str, str, int]:
        started = _now_ms()
        if content_url.startswith("data:"):
            header, separator, encoded = content_url.partition(",")
            if not separator or ";base64" not in header.lower():
                raise GatewayContractError(f"{period.upper()} export returned an unsupported inline payload.")
            try:
                content = base64.b64decode(encoded, validate=True)
            except Exception as exc:
                raise GatewayContractError(f"{period.upper()} inline workbook payload is invalid base64.") from exc
            content_type = header[5:].split(";", 1)[0] or "application/octet-stream"
            filename = f"gateway-downlines-{period}-{report_date or (self.identity.gateway_date if self.identity else 'unknown')}.xlsx"
            return content, content_type, filename, _now_ms() - started
        absolute = _require_https(urljoin(self.config.stats_origin, content_url), allowed_origin=self.config.stats_origin)
        timeout = self.config.export_timeout_seconds
        if deadline is not None:
            timeout = _remaining_timeout(deadline, timeout, f"{period.upper()} export")
        response = await self.session.get(absolute, timeout=timeout)
        _require_success_status(response, f"{period.upper()} workbook download")
        disposition = response.headers.get("content-disposition", "")
        filename_match = re.search(r'filename\*?=(?:UTF-8\'\')?["\']?([^"\';]+)', disposition, re.I)
        filename = filename_match.group(1).strip() if filename_match else (
            f"gateway-downlines-{period}-{report_date or (self.identity.gateway_date if self.identity else 'unknown')}.xlsx"
        )
        return response.content, response.headers.get("content-type", ""), filename, _now_ms() - started

    def _coverage(
        self,
        period: str,
        evidence: WorkbookEvidence,
        *,
        identity: ReportIdentity | None = None,
    ) -> dict[str, Any]:
        active_identity = identity or self.identity
        if not active_identity:
            raise GatewayContractError("StatS report identity is not ready.")
        if period == "wtd":
            current = active_identity.current_week
            return {
                "periodStart": current.period_start,
                "periodEnd": current.period_end,
                "weekStart": evidence.file_date,
                "weekOrdinal": current.week_ordinal,
                "periodFinal": current.period_final,
                "reportAsOfDate": active_identity.gateway_date,
            }
        if period == "mtd":
            return {
                "periodStart": evidence.file_date,
                "periodEnd": active_identity.current_week.period_end,
                "reportAsOfDate": active_identity.gateway_date,
            }
        return {
            "periodStart": evidence.file_date,
            "periodEnd": f"{evidence.file_date[:4]}-12-31",
            "reportAsOfDate": active_identity.gateway_date,
        }

    async def _direct_export(
        self,
        period: str,
        *,
        identity: ReportIdentity | None = None,
        report_url: str | None = None,
    ) -> GatewayArtifact:
        active_identity = identity or self.identity
        if not active_identity:
            raise GatewayContractError("StatS report identity is not ready.")
        started = _now_ms()
        response = await self.session.get(
            f"{self.config.stats_origin}/report/ip-download",
            params=self._export_params(period, identity=active_identity),
            headers={"X-Requested-With": "XMLHttpRequest", "Referer": report_url or self.report_url or self.config.stats_origin},
            timeout=self.config.export_timeout_seconds,
        )
        generated_ms = _now_ms() - started
        _require_success_status(response, f"{period.upper()} export")
        payload = _safe_json(response, f"{period.upper()} export")
        content_url = payload.get("content")
        if not isinstance(content_url, str) or not content_url:
            raise GatewayContractError(f"{period.upper()} export did not return a workbook location.")
        content, content_type, filename, download_ms = await self._fetch_workbook(
            content_url,
            period=period,
            report_date=active_identity.gateway_date,
        )
        evidence, rows = await asyncio.to_thread(
            _parse_workbook,
            content,
            period=period,
            identity=active_identity,
            content_type=content_type,
        )
        total_ms = _now_ms() - started
        self.metrics.export_ms[period] = total_ms
        return GatewayArtifact(
            period_type=period,
            source_filename=filename,
            content=content,
            evidence=evidence,
            rows=rows,
            coverage=self._coverage(period, evidence, identity=active_identity),
            generate_ms=generated_ms,
            ready_ms=generated_ms,
            download_ms=download_ms,
            total_ms=total_ms,
        )

    async def _ytd_export(
        self,
        *,
        identity: ReportIdentity | None = None,
        report_url: str | None = None,
    ) -> GatewayArtifact:
        active_identity = identity or self.identity
        if not active_identity:
            raise GatewayContractError("StatS report identity is not ready.")
        period = "ytd"
        active_report_url = report_url or self.report_url or self.config.stats_origin
        started = _now_ms()
        deadline = time.monotonic() + self.config.ytd_timeout_seconds
        response = await self.session.get(
            f"{self.config.stats_origin}/report/ip-download-ytd",
            params=self._export_params(period, identity=active_identity),
            headers={"X-Requested-With": "XMLHttpRequest", "Referer": active_report_url},
            timeout=_remaining_timeout(deadline, self.config.request_timeout_seconds, "YTD export"),
        )
        generate_ms = _now_ms() - started
        _require_success_status(response, "YTD export start")
        payload = _safe_json(response, "YTD export start")
        log_id = payload.get("log_id")
        if not log_id:
            raise GatewayContractError("YTD export did not return a generation job.")
        polls = 0
        ready_ms = 0
        try:
            while time.monotonic() < deadline:
                polls += 1
                status_response = await self.session.get(
                    f"{self.config.stats_origin}/report/ip-download-status",
                    params={"log_id": str(log_id)},
                    headers={"X-Requested-With": "XMLHttpRequest", "Referer": active_report_url},
                    timeout=_remaining_timeout(deadline, self.config.request_timeout_seconds, "YTD export"),
                )
                _require_success_status(status_response, "YTD export status")
                status_payload = _safe_json(status_response, "YTD export status")
                status = str(status_payload.get("status") or "").lower()
                if status == "completed":
                    ready_ms = _now_ms() - started
                    break
                if status in {"failed", "error", "cancelled", "canceled"}:
                    raise GatewayTransientError(f"YTD export entered terminal state: {status}.")
                base_delay = POLL_SCHEDULE_SECONDS[min(polls - 1, len(POLL_SCHEDULE_SECONDS) - 1)]
                jitter = random.uniform(-0.15, 0.15) * base_delay
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise GatewayTransientError("YTD export timed out before completion.")
                await asyncio.sleep(min(max(0.05, base_delay + jitter), remaining))
            else:
                raise GatewayTransientError("YTD export timed out before completion.")
            result_response = await self.session.get(
                f"{self.config.stats_origin}/report/YTDDownloadFile",
                params={"log_id": str(log_id)},
                headers={"X-Requested-With": "XMLHttpRequest", "Referer": active_report_url},
                timeout=_remaining_timeout(deadline, self.config.request_timeout_seconds, "YTD export"),
            )
            _require_success_status(result_response, "YTD download result")
            result_payload = _safe_json(result_response, "YTD download result")
            content_url = result_payload.get("content")
            if not isinstance(content_url, str) or not content_url:
                raise GatewayContractError("YTD completion did not return a workbook location.")
            content, content_type, filename, download_ms = await self._fetch_workbook(
                content_url,
                period=period,
                report_date=active_identity.gateway_date,
                deadline=deadline,
            )
            evidence, rows = await asyncio.to_thread(
                _parse_workbook,
                content,
                period=period,
                identity=active_identity,
                content_type=content_type,
            )
            total_ms = _now_ms() - started
            self.metrics.export_ms[period] = total_ms
            return GatewayArtifact(
                period_type=period,
                source_filename=filename,
                content=content,
                evidence=evidence,
                rows=rows,
                coverage=self._coverage(period, evidence, identity=active_identity),
                generate_ms=generate_ms,
                ready_ms=ready_ms,
                download_ms=download_ms,
                total_ms=total_ms,
                poll_count=polls,
            )
        finally:
            try:
                await self.session.get(
                    f"{self.config.stats_origin}/report/deleteYTDFileIP",
                    params={"log_id": str(log_id)},
                    headers={"X-Requested-With": "XMLHttpRequest", "Referer": active_report_url},
                    timeout=min(5.0, max(0.05, self.config.request_timeout_seconds)),
                )
            except Exception:
                pass

    def _historical_report_url(self, item: HistoricalPlanItem) -> str:
        if not self.identity:
            raise GatewayAuthenticationError("Gateway/StatS session has not been established.")
        prefix = "ipreport" if item.period_type == "wtd" else f"ipreport_{item.period_type}"
        return f"{self.config.stats_origin}/{prefix}/{self.identity.uuid}/{item.entry.date}"

    async def _historical_identity(self, item: HistoricalPlanItem) -> tuple[ReportIdentity, str]:
        if not self.identity:
            raise GatewayAuthenticationError("Gateway/StatS session has not been established.")
        report_url = self._historical_report_url(item)
        response = await self.session.get(report_url, timeout=self.config.request_timeout_seconds)
        _require_success_status(response, f"Historical {item.period_type.upper()} report page")
        if re.search(r"Page\s+Not\s+Found", response.text, re.I):
            raise GatewayTransientError(f"Gateway {item.period_type.upper()} historical report is not available yet.")
        facts = _page_facts(response.text)
        if facts.has_login_email or facts.has_login_password or facts.has_otp:
            raise GatewayAuthenticationError("Gateway/StatS browser session expired and requires reconnect.")
        if not facts.csrf:
            raise GatewayContractError("Historical StatS report page is missing its CSRF token.")
        if facts.report_uuid and facts.report_uuid != self.identity.uuid:
            raise GatewayContractError("Historical StatS report identity changed account context.")
        if facts.gateway_date and facts.gateway_date != item.entry.date:
            raise GatewayContractError("Historical StatS report page returned an unexpected weekly identity.")
        historical_identity = ReportIdentity(
            uuid=self.identity.uuid,
            gateway_date=item.entry.date,
            csrf_token=facts.csrf,
            current_week=item.entry,
            available_week_count=self.identity.available_week_count,
            oldest_week_date=self.identity.oldest_week_date,
            newest_week_date=self.identity.newest_week_date,
            weekly_entries=self.identity.weekly_entries,
        )
        return historical_identity, report_url

    async def pull_historical_artifacts(
        self,
        *,
        from_date: str,
        through_date: str,
        offset: int = 0,
        limit: int | None = None,
    ) -> HistoricalArtifactPull:
        if not self.identity:
            raise GatewayAuthenticationError("Gateway/StatS session has not been established.")
        if isinstance(offset, bool) or not isinstance(offset, int) or offset < 0:
            raise ValueError("Gateway historical offset must be a non-negative integer.")
        if limit is not None and (isinstance(limit, bool) or not isinstance(limit, int) or limit < 1):
            raise ValueError("Gateway historical limit must be a positive integer.")
        plan = historical_plan(self.identity.weekly_entries, from_date, through_date)
        if not plan:
            raise GatewayTransientError("Gateway did not expose any historical report weeks in the requested range.")
        total_count = len(plan)
        selected_plan = plan[offset:] if limit is None else plan[offset:offset + limit]
        if not selected_plan:
            return HistoricalArtifactPull(
                status="complete",
                outcomes=(),
                offset=offset,
                total_count=total_count,
            )
        started = _now_ms()
        outcomes: list[HistoricalArtifactOutcome] = []
        for index, item in enumerate(selected_plan):
            try:
                historical_identity, report_url = await self._historical_identity(item)
                artifact = (
                    await self._ytd_export(identity=historical_identity, report_url=report_url)
                    if item.period_type == "ytd"
                    else await self._direct_export(
                        item.period_type,
                        identity=historical_identity,
                        report_url=report_url,
                    )
                )
                outcomes.append(HistoricalArtifactOutcome(
                    period_type=item.period_type,
                    entry=item.entry,
                    status="validated",
                    artifact=artifact,
                ))
            except asyncio.CancelledError:
                raise
            except Exception as exc:
                error_code, retryable = classify_provider_error(exc)
                outcomes.append(HistoricalArtifactOutcome(
                    period_type=item.period_type,
                    entry=item.entry,
                    status="failed",
                    error_code=error_code,
                    retryable=retryable,
                ))
                if isinstance(exc, GatewayAuthenticationError):
                    for pending in selected_plan[index + 1:]:
                        outcomes.append(HistoricalArtifactOutcome(
                            period_type=pending.period_type,
                            entry=pending.entry,
                            status="failed",
                            error_code="authentication_failed",
                            retryable=False,
                        ))
                    break
        self.metrics.total_sync_ms = _now_ms() - started
        success_count = sum(1 for outcome in outcomes if outcome.artifact is not None)
        status = "complete" if success_count == len(outcomes) else "partial" if success_count else "failed"
        return HistoricalArtifactPull(
            status=status,
            outcomes=tuple(outcomes),
            offset=offset,
            total_count=total_count,
        )

    async def pull_current_artifacts(self) -> CurrentArtifactPull:
        if not self.identity:
            raise GatewayAuthenticationError("Gateway/StatS session has not been established.")
        started = _now_ms()
        results = await asyncio.gather(
            self._direct_export("wtd"),
            self._direct_export("mtd"),
            self._ytd_export(),
            return_exceptions=True,
        )
        self.metrics.total_sync_ms = _now_ms() - started
        outcomes: dict[str, ArtifactOutcome] = {}
        for period, result in zip(("wtd", "mtd", "ytd"), results, strict=True):
            if isinstance(result, asyncio.CancelledError):
                raise result
            if isinstance(result, BaseException):
                error_code, retryable = classify_provider_error(result)
                outcomes[period] = ArtifactOutcome(
                    period_type=period,
                    status="failed",
                    error_code=error_code,
                    retryable=retryable,
                )
            else:
                outcomes[period] = ArtifactOutcome(
                    period_type=period,
                    status="validated",
                    artifact=result,
                )
        success_count = sum(1 for outcome in outcomes.values() if outcome.artifact is not None)
        status = "complete" if success_count == 3 else "partial" if success_count else "failed"
        return CurrentArtifactPull(status=status, outcomes=outcomes)

    def diagnostic_summary(self) -> dict[str, Any]:
        """Return non-secret operational evidence suitable for structured logs."""
        return {
            "provider": "http",
            "transport": "curl_cffi/chrome",
            "gateway_login_ms": self.metrics.gateway_login_ms,
            "gateway_handoff_ms": self.metrics.gateway_handoff_ms,
            "stats_autologin_ms": self.metrics.stats_autologin_ms,
            "total_sync_ms": self.metrics.total_sync_ms,
            "export_ms": dict(self.metrics.export_ms),
        }


__all__ = [
    "ArtifactOutcome",
    "CanonicalGatewayRow",
    "CurrentArtifactPull",
    "GatewayArtifact",
    "GatewayAuthenticationError",
    "GatewayContractError",
    "GatewayHttpProvider",
    "GatewayTransientError",
    "HistoricalArtifactOutcome",
    "HistoricalArtifactPull",
    "HistoricalPlanItem",
    "HttpGatewayConfig",
    "ReportIdentity",
    "WeeklyEntry",
    "WorkbookEvidence",
    "artifact_csv_text",
    "classify_provider_error",
    "historical_plan",
    "parse_report_identity",
    "parse_stats_bootstrap",
    "validate_workbook",
]
